from __future__ import annotations

from io import BytesIO
import re
import sys
from typing import Any

from fastapi import HTTPException, status

from identity_access.application import rbac_service
from identity_access.infrastructure.persistence.common import PLATFORM_TENANT_KEY, auth_database_target, connect, require_auth_ready


IMPORT_HEADERS = ["用户名", "姓名", "邮箱", "初始密码", "角色", "部门", "主部门", "启用状态", "超级用户"]
EXPORT_HEADERS = ["ID", "用户名", "姓名", "邮箱", "角色", "部门", "主部门", "启用状态", "超级用户", "创建时间", "更新时间"]
TENANT_IMPORT_HEADERS = ["用户名", "姓名", "邮箱", "初始密码", "角色", "部门", "主部门", "启用状态"]
TENANT_EXPORT_HEADERS = ["ID", "用户名", "姓名", "邮箱", "角色", "部门", "主部门", "启用状态", "租户管理员", "创建时间", "更新时间"]
TRUE_LABELS = {"是", "启用", "true", "1", "yes", "y"}
FALSE_LABELS = {"否", "停用", "false", "0", "no", "n"}
SPLIT_PATTERN = re.compile(r"[,，;；\n]+")


def require_openpyxl() -> Any:
    try:
        import openpyxl
    except ModuleNotFoundError as exc:
        missing_name = exc.name or "openpyxl"
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail=f"Excel 导入导出依赖缺失：{missing_name}。当前后端解释器：{sys.executable}",
        ) from exc
    return openpyxl


def build_user_import_template() -> BytesIO:
    openpyxl = require_openpyxl()
    roles, departments = load_platform_options()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "用户导入"
    sheet.append(IMPORT_HEADERS)
    sheet.append(["zhangsan", "张三", "zhangsan@example.com", "ChangeMe123!", "", "", "", "启用", "否"])
    sheet.freeze_panes = "A2"
    for column_index, width in enumerate([24, 18, 28, 20, 28, 28, 28, 14, 14], start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=column_index).column_letter].width = width

    options_sheet = workbook.create_sheet("选项")
    options_sheet.sheet_state = "hidden"
    write_options(options_sheet, 1, "角色", [option["label"] for option in roles])
    write_options(options_sheet, 3, "部门", [option["label"] for option in departments])
    write_options(options_sheet, 5, "是否", ["启用", "停用", "是", "否"])

    add_dropdown(sheet, "E2:E500", "选项", "A", len(roles))
    add_dropdown(sheet, "F2:F500", "选项", "C", len(departments))
    add_dropdown(sheet, "G2:G500", "选项", "C", len(departments))
    add_dropdown(sheet, "H2:H500", "选项", "E", 4)
    add_dropdown(sheet, "I2:I500", "选项", "E", 4)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_tenant_user_import_template(tenant_id: int) -> BytesIO:
    openpyxl = require_openpyxl()
    roles, departments = load_tenant_options(tenant_id)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "成员导入"
    sheet.append(TENANT_IMPORT_HEADERS)
    sheet.append(["zhangsan", "张三", "zhangsan@example.com", "ChangeMe123!", "", "", "", "启用"])
    sheet.freeze_panes = "A2"
    for column_index, width in enumerate([24, 18, 28, 20, 28, 28, 28, 14], start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=column_index).column_letter].width = width

    options_sheet = workbook.create_sheet("选项")
    options_sheet.sheet_state = "hidden"
    write_options(options_sheet, 1, "角色", [option["label"] for option in roles])
    write_options(options_sheet, 3, "部门", [option["label"] for option in departments])
    write_options(options_sheet, 5, "状态", ["启用", "停用"])

    add_dropdown(sheet, "E2:E500", "选项", "A", len(roles))
    add_dropdown(sheet, "F2:F500", "选项", "C", len(departments))
    add_dropdown(sheet, "G2:G500", "选项", "C", len(departments))
    add_dropdown(sheet, "H2:H500", "选项", "E", 2)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_user_export_workbook(users: list[dict[str, Any]]) -> BytesIO:
    openpyxl = require_openpyxl()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "用户"
    sheet.append(EXPORT_HEADERS)
    for user in users:
        departments = list(user.get("departments") or [])
        primary = next((item for item in departments if bool(item.get("is_primary"))), None)
        sheet.append(
            [
                user.get("id"),
                user.get("username", ""),
                user.get("full_name", ""),
                user.get("email", ""),
                join_names(user.get("roles") or []),
                join_names(departments),
                str((primary or {}).get("name", "") or ""),
                "启用" if bool(user.get("is_active", True)) else "停用",
                "是" if bool(user.get("is_superuser", False)) else "否",
                user.get("create_time", ""),
                user.get("update_time", ""),
            ]
        )
    sheet.freeze_panes = "A2"
    for column_index, width in enumerate([10, 24, 18, 28, 28, 28, 28, 14, 14, 24, 24], start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=column_index).column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_tenant_user_export_workbook(users: list[dict[str, Any]]) -> BytesIO:
    openpyxl = require_openpyxl()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "成员"
    sheet.append(TENANT_EXPORT_HEADERS)
    for user in users:
        departments = list(user.get("departments") or [])
        primary = next((item for item in departments if bool(item.get("is_primary"))), None)
        sheet.append(
            [
                user.get("id"),
                user.get("username", ""),
                user.get("full_name", ""),
                user.get("email", ""),
                join_names(user.get("roles") or []),
                join_names(departments),
                str((primary or {}).get("name", "") or ""),
                "启用" if bool(user.get("is_active", True)) else "停用",
                "是" if bool(user.get("is_tenant_admin", False)) else "否",
                user.get("create_time", ""),
                user.get("update_time", ""),
            ]
        )
    sheet.freeze_panes = "A2"
    for column_index, width in enumerate([10, 24, 18, 28, 28, 28, 28, 14, 14, 24, 24], start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=column_index).column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def parse_user_import_workbook(content: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    openpyxl = require_openpyxl()
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="无法读取用户导入文件，请上传 xlsx 文件") from exc
    sheet = workbook["用户导入"] if "用户导入" in workbook.sheetnames else workbook.active
    headers = [normalize_cell(sheet.cell(row=1, column=index).value) for index in range(1, len(IMPORT_HEADERS) + 1)]
    if headers != IMPORT_HEADERS:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="导入模板表头不正确，请使用中文模板列")

    roles, departments = load_platform_options()
    role_by_label = option_lookup(roles)
    department_by_label = option_lookup(departments)
    rows: list[dict[str, Any]] = []
    assignments: list[dict[str, Any]] = []
    platform_tenant_id = get_platform_tenant_id()
    for row_index in range(2, sheet.max_row + 1):
        values = [normalize_cell(sheet.cell(row=row_index, column=index).value) for index in range(1, len(IMPORT_HEADERS) + 1)]
        if not any(values):
            continue
        role_keys = resolve_multi_options(values[4], role_by_label, "角色", row_index)
        department_ids = [int(value) for value in resolve_multi_options(values[5], department_by_label, "部门", row_index)]
        primary_department_id = resolve_single_option(values[6], department_by_label, "主部门", row_index)
        if primary_department_id and primary_department_id not in department_ids:
            department_ids.insert(0, int(primary_department_id))
        rows.append(
            {
                "tenant_id": platform_tenant_id,
                "username": values[0],
                "full_name": values[1],
                "email": values[2],
                "password": values[3],
                "role_keys": role_keys,
                "is_active": parse_bool(values[7], default=True),
                "is_superuser": parse_bool(values[8], default=False),
            }
        )
        assignments.append(
            {
                "tenant_id": platform_tenant_id,
                "department_ids": department_ids,
                "primary_department_id": int(primary_department_id) if primary_department_id else None,
            }
        )
    if not rows:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="导入文件没有用户数据")
    return rows, assignments


def parse_tenant_user_import_workbook(content: bytes, *, tenant_id: int) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    openpyxl = require_openpyxl()
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="无法读取成员导入文件，请上传 xlsx 文件") from exc
    sheet = workbook["成员导入"] if "成员导入" in workbook.sheetnames else workbook.active
    headers = [normalize_cell(sheet.cell(row=1, column=index).value) for index in range(1, len(TENANT_IMPORT_HEADERS) + 1)]
    if headers != TENANT_IMPORT_HEADERS:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="导入模板表头不正确，请使用中文模板列")

    roles, departments = load_tenant_options(tenant_id)
    role_by_label = option_lookup(roles)
    department_by_label = option_lookup(departments)
    rows: list[dict[str, Any]] = []
    assignments: list[dict[str, Any]] = []
    for row_index in range(2, sheet.max_row + 1):
        values = [normalize_cell(sheet.cell(row=row_index, column=index).value) for index in range(1, len(TENANT_IMPORT_HEADERS) + 1)]
        if not any(values):
            continue
        role_keys = resolve_multi_options(values[4], role_by_label, "角色", row_index)
        department_ids = [int(value) for value in resolve_multi_options(values[5], department_by_label, "部门", row_index)]
        primary_department_id = resolve_single_option(values[6], department_by_label, "主部门", row_index)
        if primary_department_id and primary_department_id not in department_ids:
            department_ids.insert(0, int(primary_department_id))
        rows.append(
            {
                "tenant_id": tenant_id,
                "username": values[0],
                "full_name": values[1],
                "email": values[2],
                "password": values[3],
                "role_keys": role_keys,
                "is_active": parse_bool(values[7], default=True),
                "is_superuser": False,
            }
        )
        assignments.append(
            {
                "tenant_id": tenant_id,
                "department_ids": department_ids,
                "primary_department_id": int(primary_department_id) if primary_department_id else None,
                "is_tenant_admin": "tenant-admin" in role_keys,
            }
        )
    if not rows:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="导入文件没有成员数据")
    return rows, assignments


def load_platform_options() -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    roles = [
        {"label": option_label(role, key_name="key"), "value": str(role.get("key", ""))}
        for role in rbac_service.list_roles()
        if str(role.get("role_scope", "platform") or "platform") == "platform"
    ]
    departments: list[dict[str, str]] = []
    try:
        from organization.application import services as organization_services

        for department in organization_services.list_departments(tenant_id=get_platform_tenant_id(), include_disabled=False).get("items", []):
            departments.append({"label": option_label(department, key_name="code"), "value": str(department.get("id", ""))})
    except Exception:
        departments = []
    return roles, departments


def load_tenant_options(tenant_id: int) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    roles = [
        {"label": option_label(role, key_name="key"), "value": str(role.get("key", ""))}
        for role in rbac_service.list_roles()
        if str(role.get("role_scope", "platform") or "platform") == "tenant"
    ]
    departments: list[dict[str, str]] = []
    try:
        from organization.application import services as organization_services

        for department in organization_services.list_departments(tenant_id=tenant_id, include_disabled=False).get("items", []):
            departments.append({"label": option_label(department, key_name="code"), "value": str(department.get("id", ""))})
    except Exception:
        departments = []
    return roles, departments


def get_platform_tenant_id() -> int:
    with connect(auth_database_target(), readonly=True) as conn:
        require_auth_ready(conn)
        row = conn.execute("SELECT id FROM tenants WHERE tenant_key = ?", (PLATFORM_TENANT_KEY,)).fetchone()
    if not row:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="platform tenant missing")
    return int(row["id"])


def write_options(sheet: Any, column: int, title: str, values: list[str]) -> None:
    sheet.cell(row=1, column=column, value=title)
    for index, value in enumerate(values, start=2):
        sheet.cell(row=index, column=column, value=value)


def add_dropdown(sheet: Any, cells: str, options_sheet: str, column: str, item_count: int) -> None:
    if item_count <= 0:
        return
    openpyxl = require_openpyxl()
    validation = openpyxl.worksheet.datavalidation.DataValidation(
        type="list",
        formula1=f"='{options_sheet}'!${column}$2:${column}${item_count + 1}",
        allow_blank=True,
    )
    sheet.add_data_validation(validation)
    validation.add(cells)


def option_label(item: dict[str, Any], *, key_name: str) -> str:
    name = str(item.get("name", "") or item.get("label", "") or item.get(key_name, ""))
    key = str(item.get(key_name, "") or "")
    return f"{name}({key})" if key and key != name else name


def option_lookup(options: list[dict[str, str]]) -> dict[str, str]:
    lookup: dict[str, str] = {}
    for option in options:
        label = str(option.get("label", "") or "").strip()
        value = str(option.get("value", "") or "").strip()
        if not label or not value:
            continue
        lookup[label] = value
        lookup[value] = value
        if "(" in label and label.endswith(")"):
            lookup[label.rsplit("(", 1)[0].strip()] = value
            lookup[label.rsplit("(", 1)[1][:-1].strip()] = value
    return lookup


def resolve_multi_options(value: str, lookup: dict[str, str], field_name: str, row_index: int) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for item in [part.strip() for part in SPLIT_PATTERN.split(value) if part.strip()]:
        resolved = lookup.get(item)
        if not resolved:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"第 {row_index} 行{field_name}不存在：{item}")
        if resolved not in seen:
            result.append(resolved)
            seen.add(resolved)
    return result


def resolve_single_option(value: str, lookup: dict[str, str], field_name: str, row_index: int) -> str | None:
    if not value:
        return None
    resolved = lookup.get(value.strip())
    if not resolved:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"第 {row_index} 行{field_name}不存在：{value}")
    return resolved


def parse_bool(value: str, *, default: bool) -> bool:
    normalized = value.strip().lower()
    if not normalized:
        return default
    if normalized in TRUE_LABELS:
        return True
    if normalized in FALSE_LABELS:
        return False
    return default


def join_names(items: list[dict[str, Any]]) -> str:
    return "，".join(str(item.get("name") or item.get("label") or item.get("key") or item.get("code") or "") for item in items)


def normalize_cell(value: Any) -> str:
    return str(value or "").strip()
