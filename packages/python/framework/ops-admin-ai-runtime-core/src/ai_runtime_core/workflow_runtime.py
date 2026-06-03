from __future__ import annotations

import ast
import base64
import copy
import json
import re
import time
import urllib.parse
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time as datetime_time
from io import BytesIO
from typing import Any, Callable, Iterator
from xml.etree import ElementTree

from ai_runtime_core.prompt_runtime import media_content_parts
from ai_runtime_core.prompt_runtime import render_template
from ai_runtime_core.prompt_runtime import resolve_variable_value


WorkflowDefinition = dict[str, Any]
WorkflowContext = dict[str, Any]
JSON_CODE_FENCE_PATTERN = re.compile(r"```(?:json)?\s*(.*?)\s*```", re.IGNORECASE | re.DOTALL)
MAX_SCRIPT_CHARS = 20000
MAX_SCRIPT_AST_NODES = 1000
MAX_SCRIPT_RANGE_SIZE = 10000
DEFAULT_FILE_EXTRACT_MAX_CHARS = 50000
MAX_FILE_EXTRACT_CHARS = 500000
MAX_FILE_EXTRACT_DATA_URL_CHARS = 16_000_000
DOCX_WORD_NAMESPACE = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PDF_FILE_MIME_TYPES = {"application/pdf"}
DOCX_FILE_MIME_TYPES = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
XLSX_FILE_MIME_TYPES = {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
TEXT_FILE_ENCODINGS = ("utf-8-sig", "utf-16", "gb18030", "big5", "latin-1")
TEXT_FILE_MIME_TYPES = {
    "application/csv",
    "application/json",
    "application/ld+json",
    "application/log",
    "application/markdown",
    "application/toml",
    "application/x-ndjson",
    "application/x-yaml",
    "application/xml",
    "application/yaml",
    "text/csv",
    "text/markdown",
    "text/plain",
    "text/xml",
    "text/yaml",
}
TEXT_FILE_EXTENSIONS = {
    ".csv",
    ".json",
    ".jsonl",
    ".log",
    ".md",
    ".ndjson",
    ".text",
    ".toml",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}
DISALLOWED_SCRIPT_NODE_TYPES = (
    ast.AsyncFor,
    ast.AsyncFunctionDef,
    ast.AsyncWith,
    ast.Await,
    ast.ClassDef,
    ast.Delete,
    ast.FunctionDef,
    ast.Global,
    ast.Import,
    ast.ImportFrom,
    ast.Lambda,
    ast.Nonlocal,
    ast.Raise,
    ast.Try,
    ast.While,
    ast.With,
    ast.Yield,
    ast.YieldFrom,
)
DISALLOWED_SCRIPT_NAMES = {
    "__builtins__",
    "__import__",
    "breakpoint",
    "compile",
    "delattr",
    "dir",
    "eval",
    "exec",
    "exit",
    "getattr",
    "globals",
    "help",
    "locals",
    "memoryview",
    "object",
    "open",
    "quit",
    "setattr",
    "super",
    "type",
    "vars",
}
DISALLOWED_SCRIPT_ATTRIBUTES = {"mro", "subclasses"}


@dataclass(slots=True)
class WorkflowLLMRequest:
    node_id: str
    model: str
    messages: list[dict[str, Any]]
    temperature: float | None = None
    response_format: dict[str, Any] | None = None
    extra_body: dict[str, Any] = field(default_factory=dict)


@dataclass(slots=True)
class WorkflowLLMResult:
    answer: Any
    usage: dict[str, Any] = field(default_factory=dict)
    model: str = ""


@dataclass(slots=True)
class WorkflowSQLRequest:
    node_id: str
    sql: str
    params: list[Any] = field(default_factory=list)
    output_key: str = ""
    result_shape: str = "rows"
    max_rows: int = 100
    data_access: dict[str, Any] = field(default_factory=dict)


@dataclass(slots=True)
class WorkflowSQLResult:
    rows: list[dict[str, Any]] = field(default_factory=list)
    columns: list[str] = field(default_factory=list)
    row_count: int = 0
    truncated: bool = False


@dataclass(slots=True)
class WorkflowFileExtractRequest:
    node_id: str
    value: Any
    output_key: str = ""
    max_chars: int = DEFAULT_FILE_EXTRACT_MAX_CHARS


@dataclass(slots=True)
class WorkflowFileExtractResult:
    text: str = ""
    files: list[dict[str, Any]] = field(default_factory=list)
    file_count: int = 0
    truncated: bool = False


@dataclass(slots=True)
class WorkflowRunResult:
    answer: str
    context: WorkflowContext
    trace: dict[str, Any]
    usage: dict[str, Any]


LLMExecutor = Callable[[WorkflowLLMRequest], WorkflowLLMResult]
SQLExecutor = Callable[[WorkflowSQLRequest], WorkflowSQLResult]
FileExtractor = Callable[[WorkflowFileExtractRequest], WorkflowFileExtractResult | dict[str, Any] | str]


class WorkflowRuntimeError(RuntimeError):
    pass


def execute_workflow(
    definition: WorkflowDefinition,
    variables: dict[str, Any],
    *,
    llm_executor: LLMExecutor,
    sql_executor: SQLExecutor | None = None,
    file_extractor: FileExtractor | None = None,
    max_steps: int = 50,
) -> WorkflowRunResult:
    result: WorkflowRunResult | None = None
    for event in iter_workflow_events(
        definition,
        variables,
        llm_executor=llm_executor,
        sql_executor=sql_executor,
        file_extractor=file_extractor,
        max_steps=max_steps,
    ):
        if event.get("event") == "workflow.completed":
            candidate = event.get("result")
            if isinstance(candidate, WorkflowRunResult):
                result = candidate
    if result is None:
        raise WorkflowRuntimeError("workflow did not complete")
    return result


def iter_workflow_events(
    definition: WorkflowDefinition,
    variables: dict[str, Any],
    *,
    llm_executor: LLMExecutor,
    sql_executor: SQLExecutor | None = None,
    file_extractor: FileExtractor | None = None,
    max_steps: int = 50,
) -> Iterator[dict[str, Any]]:
    normalized = normalize_workflow_definition(definition)
    nodes = {node["id"]: node for node in normalized["nodes"]}
    outgoing = outgoing_edges(normalized["edges"])
    current_id = start_node_id(normalized)
    context: WorkflowContext = {
        "variables": dict(variables),
        "nodes": {},
        "last": {},
    }
    trace_nodes: list[dict[str, Any]] = []
    usage: dict[str, Any] = {}
    answer = ""

    for _ in range(max_steps):
        node = nodes.get(current_id)
        if not node:
            raise WorkflowRuntimeError(f"workflow node not found: {current_id}")
        node_type = str(node.get("type") or "").strip()
        started_at = time.perf_counter()
        trace_node: dict[str, Any] = {
            "node_id": current_id,
            "node_type": node_type,
            "status": "running",
        }
        yield {"event": "workflow.node.started", "node": snapshot_value(trace_node)}
        try:
            if node_type == "start":
                output = {"variables": snapshot_value(context["variables"])}
            elif node_type == "llm":
                result = execute_llm_node(node, context, llm_executor)
                output = {"answer": result.answer, "usage": result.usage, "model": result.model}
                answer = workflow_answer_text(result.answer)
                merge_usage(usage, result.usage)
            elif node_type in {"sql", "sql_query"}:
                output = execute_sql_query_node(node, context, sql_executor)
            elif node_type in {"file_extract", "file_extraction"}:
                output = execute_file_extract_node(node, context, file_extractor)
            elif node_type in {"script", "python_script"}:
                output = execute_script_node(node, context)
            elif node_type in {"condition", "if_else"}:
                matched = evaluate_condition_node(node, context)
                output = {"matched": matched}
                trace_node["branch"] = "true" if matched else "false"
            elif node_type == "end":
                output = execute_end_node(node, context, answer)
                answer = str(output.get("answer") or answer)
                trace_node["status"] = "success"
                trace_node["output"] = snapshot_value(output)
                trace_node["elapsed_ms"] = elapsed_ms(started_at)
                trace_nodes.append(trace_node)
                yield {"event": "workflow.node.completed", "node": snapshot_value(trace_node)}
                break
            else:
                raise WorkflowRuntimeError(f"unsupported workflow node type: {node_type}")
            context["nodes"][current_id] = output
            context["last"] = output
            trace_node["status"] = "success"
            trace_node["output"] = snapshot_value(output)
            trace_node["elapsed_ms"] = elapsed_ms(started_at)
            trace_nodes.append(trace_node)
            yield {"event": "workflow.node.completed", "node": snapshot_value(trace_node)}
            next_id = next_node_id(current_id, outgoing, trace_node.get("branch"))
            if not next_id:
                break
            current_id = next_id
        except Exception as exc:
            trace_node["status"] = "failed"
            trace_node["error"] = str(exc)
            trace_node["elapsed_ms"] = elapsed_ms(started_at)
            trace_nodes.append(trace_node)
            yield {"event": "workflow.node.failed", "node": snapshot_value(trace_node)}
            raise
    else:
        raise WorkflowRuntimeError("workflow exceeded max steps")

    yield {
        "event": "workflow.completed",
        "result": WorkflowRunResult(
            answer=answer,
            context=context,
            usage=usage,
            trace={
                "workflow": {
                    "nodes": trace_nodes,
                    "final_node_id": current_id,
                }
            },
        ),
    }


def snapshot_value(value: Any) -> Any:
    return copy.deepcopy(value)


def normalize_workflow_definition(definition: WorkflowDefinition) -> WorkflowDefinition:
    nodes = definition.get("nodes") if isinstance(definition.get("nodes"), list) else []
    edges = definition.get("edges") if isinstance(definition.get("edges"), list) else []
    normalized_nodes = [normalize_node(item) for item in nodes if isinstance(item, dict)]
    normalized_edges = [normalize_edge(item) for item in edges if isinstance(item, dict)]
    if not normalized_nodes:
        raise WorkflowRuntimeError("workflow requires at least one node")
    if not any(node["type"] == "start" for node in normalized_nodes):
        raise WorkflowRuntimeError("workflow requires a start node")
    return {"nodes": normalized_nodes, "edges": normalized_edges}


def normalize_node(node: dict[str, Any]) -> dict[str, Any]:
    node_id = str(node.get("id") or "").strip()
    node_type = str(node.get("type") or node.get("node_type") or "").strip()
    if not node_id:
        raise WorkflowRuntimeError("workflow node id is required")
    if not node_type:
        raise WorkflowRuntimeError(f"workflow node type is required: {node_id}")
    data = node.get("data") if isinstance(node.get("data"), dict) else {}
    return {"id": node_id, "type": node_type, "data": data}


def normalize_edge(edge: dict[str, Any]) -> dict[str, Any]:
    source = str(edge.get("source") or "").strip()
    target = str(edge.get("target") or "").strip()
    if not source or not target:
        raise WorkflowRuntimeError("workflow edge source and target are required")
    return {
        "id": str(edge.get("id") or f"{source}-{target}"),
        "source": source,
        "target": target,
        "source_handle": str(edge.get("sourceHandle") or edge.get("source_handle") or "").strip(),
    }


def start_node_id(definition: WorkflowDefinition) -> str:
    for node in definition["nodes"]:
        if node["type"] == "start":
            return str(node["id"])
    raise WorkflowRuntimeError("workflow requires a start node")


def outgoing_edges(edges: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {}
    for edge in edges:
        result.setdefault(str(edge["source"]), []).append(edge)
    return result


def next_node_id(current_id: str, outgoing: dict[str, list[dict[str, Any]]], branch: Any = None) -> str:
    edges = outgoing.get(current_id) or []
    if branch in {"true", "false"}:
        for edge in edges:
            if edge.get("source_handle") == branch:
                return str(edge["target"])
    return str(edges[0]["target"]) if edges else ""


def execute_llm_node(node: dict[str, Any], context: WorkflowContext, llm_executor: LLMExecutor) -> WorkflowLLMResult:
    data = node["data"]
    model = str(data.get("model") or data.get("route_key") or "").strip()
    if not model:
        raise WorkflowRuntimeError(f"LLM node model is required: {node['id']}")
    variables = workflow_template_context(context)
    messages = render_llm_messages(data, variables)
    if not messages:
        raise WorkflowRuntimeError(f"LLM node prompt is required: {node['id']}")
    result = llm_executor(
        WorkflowLLMRequest(
            node_id=node["id"],
            model=model,
            messages=messages,
            temperature=float(data["temperature"]) if data.get("temperature") is not None else None,
            response_format=data.get("response_format") if isinstance(data.get("response_format"), dict) else None,
            extra_body=data.get("extra_body") if isinstance(data.get("extra_body"), dict) else {},
        )
    )
    result = maybe_decode_json_response(result, data)
    output_key = str(data.get("output_key") or "").strip()
    if output_key:
        assign_path(context["variables"], output_key, result.answer)
    return result


def execute_sql_query_node(
    node: dict[str, Any],
    context: WorkflowContext,
    sql_executor: SQLExecutor | None,
) -> dict[str, Any]:
    if sql_executor is None:
        raise WorkflowRuntimeError(f"SQL executor is required for node: {node['id']}")
    data = node["data"]
    variables = workflow_template_context(context)
    sql = render_template(str(data.get("sql") or data.get("query") or ""), variables).strip()
    if not sql:
        raise WorkflowRuntimeError(f"SQL query is required: {node['id']}")
    if uses_named_sql_params(sql):
        sql, params = expand_named_sql_params(sql, render_named_sql_params(data.get("params"), variables))
    else:
        params = render_sql_params(data.get("params"), variables)
    ensure_readonly_sql(sql)
    output_key = str(data.get("output_key") or "").strip()
    if not output_key:
        raise WorkflowRuntimeError(f"SQL node output_key is required: {node['id']}")
    result_shape = str(data.get("result_shape") or "rows").strip().lower()
    if result_shape not in {"rows", "first", "scalar"}:
        raise WorkflowRuntimeError(f"unsupported SQL result_shape: {result_shape}")
    max_rows = bounded_max_rows(data.get("max_rows"))
    result = sql_executor(
        WorkflowSQLRequest(
            node_id=node["id"],
            sql=sql,
            params=params,
            output_key=output_key,
            result_shape=result_shape,
            max_rows=max_rows,
            data_access=data.get("data_access") if isinstance(data.get("data_access"), dict) else {},
        )
    )
    output = sql_output_payload(result, result_shape)
    assign_path(context["variables"], output_key, output)
    return output


def render_sql_params(value: Any, variables: dict[str, Any]) -> list[Any]:
    if value is None or value == "":
        return []
    if isinstance(value, str):
        rendered = render_template(value, variables).strip()
        if not rendered:
            return []
        try:
            value = json.loads(rendered)
        except json.JSONDecodeError:
            return [rendered]
    if isinstance(value, list):
        return [resolve_sql_param(item, variables) for item in value]
    if isinstance(value, dict):
        return [resolve_sql_param(value[key], variables) for key in sorted(value)]
    raise WorkflowRuntimeError("SQL params must be an array, object, JSON string, or empty")


def uses_named_sql_params(sql: str) -> bool:
    return re.search(r"(?<!:):[a-zA-Z_][a-zA-Z0-9_]*", sql) is not None


def render_named_sql_params(value: Any, variables: dict[str, Any]) -> dict[str, Any]:
    if value is None or value == "":
        return {}
    if isinstance(value, str):
        rendered = render_template(value, variables).strip()
        if not rendered:
            return {}
        try:
            value = json.loads(rendered)
        except json.JSONDecodeError as exc:
            raise WorkflowRuntimeError(f"SQL named params JSON is invalid: {exc}") from exc
    if not isinstance(value, dict):
        raise WorkflowRuntimeError("SQL named params must be a JSON object when SQL uses :name placeholders")
    return {str(key): resolve_sql_param(param_value, variables) for key, param_value in value.items()}


def expand_named_sql_params(sql: str, params: dict[str, Any]) -> tuple[str, list[Any]]:
    values: list[Any] = []

    def replace(match: re.Match[str]) -> str:
        name = match.group(0)[1:]
        if name not in params:
            raise WorkflowRuntimeError(f"SQL named param is missing: {name}")
        value = params[name]
        if isinstance(value, (list, tuple)):
            if not value:
                return "NULL"
            values.extend(value)
            return ", ".join("?" for _ in value)
        values.append(value)
        return "?"

    return re.sub(r"(?<!:):[a-zA-Z_][a-zA-Z0-9_]*", replace, sql), values


def resolve_sql_param(value: Any, variables: dict[str, Any]) -> Any:
    if isinstance(value, str):
        stripped = value.strip()
        if stripped.startswith("{{") and stripped.endswith("}}"):
            return resolve_variable_value(variables, stripped[2:-2].strip())
        return render_template(value, variables)
    return value


def ensure_readonly_sql(sql: str) -> None:
    compact = strip_sql_comments(sql).strip()
    if not compact:
        raise WorkflowRuntimeError("SQL query is required")
    statements = [item.strip() for item in compact.split(";") if item.strip()]
    if len(statements) > 1 or (compact.endswith(";") and len(statements) != 1):
        raise WorkflowRuntimeError("SQL node only supports a single read-only statement")
    first_token = first_sql_token(compact)
    if first_token not in {"select", "with"}:
        raise WorkflowRuntimeError("SQL node only supports SELECT queries")
    dangerous = re.search(
        r"\b(insert|update|delete|drop|alter|truncate|create|replace|merge|call|execute|grant|revoke)\b",
        compact,
        flags=re.IGNORECASE,
    )
    if dangerous:
        raise WorkflowRuntimeError(f"SQL node rejected non-read-only keyword: {dangerous.group(1).upper()}")


def strip_sql_comments(sql: str) -> str:
    without_line_comments = re.sub(r"--.*?(?=\r?\n|$)", " ", sql)
    return re.sub(r"/\*.*?\*/", " ", without_line_comments, flags=re.DOTALL)


def first_sql_token(sql: str) -> str:
    match = re.match(r"\s*([a-zA-Z_]+)", sql)
    return match.group(1).lower() if match else ""


def bounded_max_rows(value: Any) -> int:
    try:
        parsed = int(value or 100)
    except (TypeError, ValueError):
        parsed = 100
    return max(1, min(1000, parsed))


def sql_output_payload(result: WorkflowSQLResult, result_shape: str) -> dict[str, Any]:
    rows = snapshot_value(result.rows)
    first = rows[0] if rows else None
    scalar = first.get(result.columns[0]) if isinstance(first, dict) and result.columns else None
    return {
        "rows": rows,
        "first": first,
        "scalar": scalar,
        "columns": list(result.columns),
        "row_count": int(result.row_count),
        "result_shape": result_shape,
        "truncated": bool(result.truncated),
    }


def execute_file_extract_node(
    node: dict[str, Any],
    context: WorkflowContext,
    file_extractor: FileExtractor | None,
) -> dict[str, Any]:
    data = node["data"]
    output_key = str(data.get("output_key") or "").strip()
    if not output_key:
        raise WorkflowRuntimeError(f"file extract node output_key is required: {node['id']}")
    max_chars = bounded_file_extract_max_chars(data.get("max_chars"))
    source_value = resolve_workflow_value(data.get("input") if "input" in data else "{{last}}", context)
    if file_extractor is None:
        output = extract_file_content_payload(source_value, max_chars=max_chars)
    else:
        output = normalize_file_extract_output(
            file_extractor(
                WorkflowFileExtractRequest(
                    node_id=node["id"],
                    value=snapshot_value(source_value),
                    output_key=output_key,
                    max_chars=max_chars,
                )
            ),
            max_chars=max_chars,
        )
    assign_path(context["variables"], output_key, output)
    return output


def bounded_file_extract_max_chars(value: Any) -> int:
    try:
        parsed = int(value or DEFAULT_FILE_EXTRACT_MAX_CHARS)
    except (TypeError, ValueError):
        parsed = DEFAULT_FILE_EXTRACT_MAX_CHARS
    return max(1, min(MAX_FILE_EXTRACT_CHARS, parsed))


def normalize_file_extract_output(value: WorkflowFileExtractResult | dict[str, Any] | str, *, max_chars: int) -> dict[str, Any]:
    if isinstance(value, WorkflowFileExtractResult):
        text, text_truncated = truncate_text(value.text, max_chars)
        files, files_truncated = normalize_file_extract_files(value.files, max_chars=max_chars)
        file_count = int(value.file_count or len(files))
        return {
            "text": text,
            "files": files,
            "file_count": file_count,
            "truncated": bool(value.truncated or text_truncated or files_truncated),
        }
    if isinstance(value, dict):
        text, text_truncated = truncate_text(str(value.get("text") or ""), max_chars)
        files, files_truncated = normalize_file_extract_files(
            value.get("files") if isinstance(value.get("files"), list) else [],
            max_chars=max_chars,
        )
        return {
            "text": text,
            "files": files,
            "file_count": int(value.get("file_count") or len(files)),
            "truncated": bool(value.get("truncated") or text_truncated or files_truncated),
        }
    text, text_truncated = truncate_text(str(value or ""), max_chars)
    return {"text": text, "files": [], "file_count": 0, "truncated": text_truncated}


def normalize_file_extract_files(files: list[Any], *, max_chars: int) -> tuple[list[dict[str, Any]], bool]:
    normalized: list[dict[str, Any]] = []
    truncated = False
    for item in files:
        record = snapshot_value(item) if isinstance(item, dict) else file_extract_item_metadata(item)
        if not isinstance(record, dict):
            record = {}
        file_text, file_text_truncated = truncate_text(str(record.get("text") or ""), max_chars)
        record["text"] = file_text
        record["truncated"] = bool(record.get("truncated") or file_text_truncated)
        truncated = truncated or bool(record["truncated"])
        normalized.append(record)
    return normalized, truncated


def extract_file_content_payload(value: Any, *, max_chars: int) -> dict[str, Any]:
    items = file_extract_items(value)
    if not items:
        return {"text": "", "files": [], "file_count": 0, "truncated": False}
    files: list[dict[str, Any]] = []
    text_parts: list[str] = []
    truncated = False
    remaining_chars = max_chars
    for item in items:
        if remaining_chars <= 0:
            remaining_text = file_extract_item_text(item)
            files.append({**file_extract_item_metadata(item), "text": "", "truncated": bool(remaining_text)})
            truncated = truncated or bool(remaining_text)
            continue
        file_payload = file_extract_item_payload(item, max_chars=max(remaining_chars, 1))
        files.append(file_payload)
        file_text = str(file_payload.get("text") or "")
        if file_text and remaining_chars > 0:
            kept, text_truncated = truncate_text(file_text, remaining_chars)
            text_parts.append(kept)
            remaining_chars -= len(kept)
            truncated = truncated or text_truncated or bool(file_payload.get("truncated"))
        elif file_text:
            truncated = True
        else:
            truncated = truncated or bool(file_payload.get("truncated"))
    joined = "\n\n".join(part for part in text_parts if part)
    text, text_truncated = truncate_text(joined, max_chars)
    return {
        "text": text,
        "files": files,
        "file_count": len(files),
        "truncated": bool(truncated or text_truncated),
    }


def file_extract_items(value: Any) -> list[Any]:
    if value in (None, ""):
        return []
    if isinstance(value, list):
        return [item for item in value if item not in (None, "")]
    if isinstance(value, dict):
        if isinstance(value.get("files"), list):
            return value["files"]
        if isinstance(value.get("items"), list) and not media_like_file_record(value):
            return value["items"]
    return [value]


def media_like_file_record(value: dict[str, Any]) -> bool:
    return any(key in value for key in ("type", "name", "mime_type", "data_url", "text", "file_ref", "file_id"))


def file_extract_item_payload(value: Any, *, max_chars: int) -> dict[str, Any]:
    metadata = file_extract_item_metadata(value)
    text = file_extract_item_text(value)
    kept, truncated = truncate_text(text, max_chars)
    return {
        **metadata,
        "text": kept,
        "truncated": truncated,
    }


def file_extract_item_metadata(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict):
        return {"name": "", "mime_type": "", "size": None}
    return {
        "name": str(value.get("name") or value.get("filename") or ""),
        "mime_type": str(value.get("mime_type") or value.get("content_type") or ""),
        "size": value.get("size"),
        "file_ref": value.get("file_ref"),
        "file_id": value.get("file_id"),
    }


def file_extract_item_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float, bool)):
        return str(value)
    if not isinstance(value, dict):
        return ""
    data_url = str(value.get("data_url") or "").strip()
    if data_url:
        extracted = extracted_text_from_data_url(
            data_url,
            name=str(value.get("name") or value.get("filename") or ""),
            mime_type=str(value.get("mime_type") or value.get("content_type") or ""),
        )
        if extracted:
            return extracted
    for key in ("text", "content"):
        if value.get(key) not in (None, ""):
            return str(value.get(key) or "")
    return ""


def extracted_text_from_data_url(data_url: str, *, name: str, mime_type: str) -> str:
    decoded = decode_file_data_url(data_url)
    if decoded is None:
        return ""
    data_mime_type, raw_bytes = decoded
    effective_mime_type = data_mime_type or mime_type
    if is_pdf_file(name=name, mime_type=effective_mime_type):
        return extract_pdf_text(raw_bytes)
    if is_docx_file(name=name, mime_type=effective_mime_type):
        return extract_docx_text(raw_bytes)
    if is_xlsx_file(name=name, mime_type=effective_mime_type):
        return extract_xlsx_text(raw_bytes)
    if is_text_like_file(name=name, mime_type=effective_mime_type):
        return decode_text_bytes(raw_bytes)
    return ""


def decode_file_data_url(data_url: str) -> tuple[str, bytes] | None:
    if len(data_url) > MAX_FILE_EXTRACT_DATA_URL_CHARS or not data_url.startswith("data:") or "," not in data_url:
        return None
    header, payload = data_url.split(",", 1)
    header_parts = [item.strip().lower() for item in header[5:].split(";")]
    data_mime_type = header_parts[0] if header_parts else ""
    try:
        raw_bytes = base64.b64decode(payload, validate=True) if "base64" in header_parts else urllib.parse.unquote_to_bytes(payload)
    except (ValueError, TypeError):
        return None
    return data_mime_type, raw_bytes


def extract_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        try:
            from PyPDF2 import PdfReader
        except ImportError as fallback_exc:
            raise WorkflowRuntimeError("PDF 解析需要安装 pypdf") from fallback_exc
    try:
        reader = PdfReader(BytesIO(content))
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
    except Exception as exc:
        raise WorkflowRuntimeError(f"PDF 文件解析失败: {exc}") from exc


def extract_docx_text(content: bytes) -> str:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            document_xml = archive.read("word/document.xml")
    except KeyError as exc:
        raise WorkflowRuntimeError("DOCX 文件缺少 word/document.xml") from exc
    except zipfile.BadZipFile as exc:
        raise WorkflowRuntimeError("DOCX 文件格式无效") from exc
    try:
        root = ElementTree.fromstring(document_xml)
    except ElementTree.ParseError as exc:
        raise WorkflowRuntimeError("DOCX 文件内容解析失败") from exc
    paragraphs: list[str] = []
    for paragraph in root.iter(f"{DOCX_WORD_NAMESPACE}p"):
        text_parts: list[str] = []
        for element in paragraph.iter():
            if element.tag == f"{DOCX_WORD_NAMESPACE}t" and element.text:
                text_parts.append(element.text)
            elif element.tag == f"{DOCX_WORD_NAMESPACE}tab":
                text_parts.append("\t")
            elif element.tag == f"{DOCX_WORD_NAMESPACE}br":
                text_parts.append("\n")
        paragraph_text = "".join(text_parts).strip()
        if paragraph_text:
            paragraphs.append(paragraph_text)
    return "\n".join(paragraphs)


def extract_xlsx_text(content: bytes) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise WorkflowRuntimeError("XLSX 解析需要安装 openpyxl") from exc
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise WorkflowRuntimeError(f"XLSX 文件解析失败: {exc}") from exc
    try:
        sheet_texts: list[str] = []
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                cells = [xlsx_cell_text(cell) for cell in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append("\t".join(cells))
            if rows:
                sheet_texts.append(f"## {worksheet.title}\n" + "\n".join(rows))
        return "\n\n".join(sheet_texts)
    finally:
        workbook.close()


def xlsx_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, datetime_time)):
        return value.isoformat()
    return str(value)


def decode_text_bytes(content: bytes) -> str:
    if not content:
        return ""
    candidates: list[tuple[int, str]] = []
    for encoding in TEXT_FILE_ENCODINGS:
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        candidates.append((text_mojibake_score(text), text))
    if not candidates:
        return content.decode("utf-8", errors="replace")
    return min(candidates, key=lambda item: item[0])[1]


def text_mojibake_score(text: str) -> int:
    if not text:
        return 0
    replacement_count = text.count("\ufffd")
    control_count = sum(1 for char in text if ord(char) < 32 and char not in "\t\r\n")
    mojibake_markers = ("锟", "绱", "Â", "Ã", "�")
    marker_count = sum(text.count(marker) for marker in mojibake_markers)
    cjk_count = sum(1 for char in text if "\u4e00" <= char <= "\u9fff")
    ascii_count = sum(1 for char in text if "\x20" <= char <= "\x7e")
    return replacement_count * 100 + control_count * 20 + marker_count * 8 - cjk_count - ascii_count // 10


def is_pdf_file(*, name: str, mime_type: str) -> bool:
    return mime_type.strip().lower() in PDF_FILE_MIME_TYPES or name.strip().lower().endswith(".pdf")


def is_docx_file(*, name: str, mime_type: str) -> bool:
    return mime_type.strip().lower() in DOCX_FILE_MIME_TYPES or name.strip().lower().endswith(".docx")


def is_xlsx_file(*, name: str, mime_type: str) -> bool:
    return mime_type.strip().lower() in XLSX_FILE_MIME_TYPES or name.strip().lower().endswith(".xlsx")


def is_text_like_file(*, name: str, mime_type: str) -> bool:
    normalized_mime = mime_type.strip().lower()
    if normalized_mime.startswith("text/") or normalized_mime in TEXT_FILE_MIME_TYPES:
        return True
    lower_name = name.strip().lower()
    return any(lower_name.endswith(extension) for extension in TEXT_FILE_EXTENSIONS)


def truncate_text(value: str, max_chars: int) -> tuple[str, bool]:
    text = str(value or "")
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars], True


def execute_script_node(node: dict[str, Any], context: WorkflowContext) -> dict[str, Any]:
    data = node["data"]
    code = str(data.get("code") or data.get("script") or "").strip()
    if not code:
        raise WorkflowRuntimeError(f"script code is required: {node['id']}")
    output_key = str(data.get("output_key") or "").strip()
    if not output_key:
        raise WorkflowRuntimeError(f"script node output_key is required: {node['id']}")
    validate_script_code(code)
    script_input = resolve_workflow_value(data.get("input") if "input" in data else "{{last}}", context)
    script_scope: dict[str, Any] = {
        "__builtins__": SAFE_SCRIPT_BUILTINS,
        "json": SAFE_SCRIPT_JSON,
        "input": snapshot_value(script_input),
        "variables": snapshot_value(context.get("variables") if isinstance(context.get("variables"), dict) else {}),
        "nodes": snapshot_value(context.get("nodes") if isinstance(context.get("nodes"), dict) else {}),
        "last": snapshot_value(context.get("last") if isinstance(context.get("last"), dict) else {}),
        "result": None,
        "output": None,
    }
    try:
        compiled = compile(code, f"<workflow-script:{node['id']}>", "exec")
        exec(compiled, script_scope, script_scope)
    except WorkflowRuntimeError:
        raise
    except Exception as exc:
        raise WorkflowRuntimeError(f"script node failed: {exc}") from exc
    result = script_scope.get("result")
    if result is None and script_scope.get("output") is not None:
        result = script_scope.get("output")
    if result is None:
        raise WorkflowRuntimeError(f"script node must assign result or output: {node['id']}")
    assign_path(context["variables"], output_key, snapshot_value(result))
    return {"result": snapshot_value(result), "output_key": output_key}


def validate_script_code(code: str) -> None:
    if len(code) > MAX_SCRIPT_CHARS:
        raise WorkflowRuntimeError(f"script code is too long; max {MAX_SCRIPT_CHARS} characters")
    try:
        tree = ast.parse(code, mode="exec")
    except SyntaxError as exc:
        raise WorkflowRuntimeError(f"script syntax is invalid: {exc.msg}") from exc
    nodes = list(ast.walk(tree))
    if len(nodes) > MAX_SCRIPT_AST_NODES:
        raise WorkflowRuntimeError(f"script is too complex; max {MAX_SCRIPT_AST_NODES} syntax nodes")
    for item in nodes:
        if isinstance(item, DISALLOWED_SCRIPT_NODE_TYPES):
            raise WorkflowRuntimeError(f"script statement is not allowed: {item.__class__.__name__}")
        if isinstance(item, ast.Name) and is_disallowed_script_name(item.id):
            raise WorkflowRuntimeError(f"script name is not allowed: {item.id}")
        if isinstance(item, ast.Attribute) and is_disallowed_script_attribute(item.attr):
            raise WorkflowRuntimeError(f"script attribute is not allowed: {item.attr}")


def is_disallowed_script_name(name: str) -> bool:
    return name.startswith("__") or name in DISALLOWED_SCRIPT_NAMES


def is_disallowed_script_attribute(name: str) -> bool:
    return name.startswith("_") or name in DISALLOWED_SCRIPT_ATTRIBUTES


def safe_script_range(*args: int) -> range:
    if not 1 <= len(args) <= 3:
        raise WorkflowRuntimeError("range expects 1 to 3 integer arguments")
    values = [int(item) for item in args]
    range_value = range(*values)
    try:
        range_size = len(range_value)
    except OverflowError as exc:
        raise WorkflowRuntimeError(f"range is too large; max {MAX_SCRIPT_RANGE_SIZE} items") from exc
    if range_size > MAX_SCRIPT_RANGE_SIZE:
        raise WorkflowRuntimeError(f"range is too large; max {MAX_SCRIPT_RANGE_SIZE} items")
    return range_value


class SafeScriptJson:
    @staticmethod
    def loads(value: str) -> Any:
        return json.loads(value)

    @staticmethod
    def dumps(value: Any, *, ensure_ascii: bool = False, indent: int | None = None) -> str:
        return json.dumps(value, ensure_ascii=ensure_ascii, indent=indent)


SAFE_SCRIPT_JSON = SafeScriptJson()
SAFE_SCRIPT_BUILTINS: dict[str, Any] = {
    "abs": abs,
    "all": all,
    "any": any,
    "bool": bool,
    "dict": dict,
    "enumerate": enumerate,
    "filter": filter,
    "float": float,
    "int": int,
    "isinstance": isinstance,
    "len": len,
    "list": list,
    "map": map,
    "max": max,
    "min": min,
    "range": safe_script_range,
    "round": round,
    "set": set,
    "sorted": sorted,
    "str": str,
    "sum": sum,
    "tuple": tuple,
    "zip": zip,
}


def render_llm_messages(data: dict[str, Any], variables: dict[str, Any]) -> list[dict[str, Any]]:
    messages: list[dict[str, Any]] = []
    for role, field_name in (
        ("system", "system_prompt"),
        ("developer", "developer_prompt"),
        ("user", "user_prompt_template"),
    ):
        content = render_template(str(data.get(field_name) or ""), variables)
        if not content.strip():
            continue
        messages.append({"role": role, "content": render_user_content(content, variables) if role == "user" else content})
    return messages


def render_user_content(text: str, variables: dict[str, Any]) -> str | list[dict[str, Any]]:
    media_parts = media_content_parts(variables)
    return [{"type": "text", "text": text}, *media_parts] if media_parts else text


def evaluate_condition_node(node: dict[str, Any], context: WorkflowContext) -> bool:
    data = node["data"]
    conditions = data.get("conditions") if isinstance(data.get("conditions"), list) else []
    if not conditions:
        conditions = [{"left": data.get("left"), "operator": data.get("operator") or "exists", "right": data.get("right")}]
    mode = str(data.get("mode") or "all").strip().lower()
    results = [evaluate_condition(item, context) for item in conditions if isinstance(item, dict)]
    return any(results) if mode == "any" else all(results)


def evaluate_condition(condition: dict[str, Any], context: WorkflowContext) -> bool:
    left = resolve_workflow_value(condition.get("left"), context)
    right = resolve_workflow_value(condition.get("right"), context)
    operator = str(condition.get("operator") or "exists").strip()
    if operator == "exists":
        return left not in (None, "")
    if operator == "empty":
        return left in (None, "")
    if operator == "equals":
        return str(left) == str(right)
    if operator == "not_equals":
        return str(left) != str(right)
    if operator == "contains":
        return str(right) in str(left)
    if operator == "not_contains":
        return str(right) not in str(left)
    if operator in {"gt", "gte", "lt", "lte"}:
        left_number = float(left or 0)
        right_number = float(right or 0)
        return {
            "gt": left_number > right_number,
            "gte": left_number >= right_number,
            "lt": left_number < right_number,
            "lte": left_number <= right_number,
        }[operator]
    raise WorkflowRuntimeError(f"unsupported condition operator: {operator}")


def execute_end_node(node: dict[str, Any], context: WorkflowContext, answer: str) -> dict[str, Any]:
    data = node["data"]
    output = data.get("output")
    if output is None:
        output = data.get("output_template")
    if output is None:
        return {"answer": answer}
    resolved = resolve_workflow_value(output, context)
    return {"answer": str(resolved or "")}


def resolve_workflow_value(value: Any, context: WorkflowContext) -> Any:
    if not isinstance(value, str):
        return value
    variables = workflow_template_context(context)
    stripped = value.strip()
    if stripped.startswith("{{") and stripped.endswith("}}"):
        return resolve_variable_value(variables, stripped[2:-2].strip())
    return render_template(value, variables)


def workflow_template_context(context: WorkflowContext) -> dict[str, Any]:
    variables = context.get("variables") if isinstance(context.get("variables"), dict) else {}
    nodes = context.get("nodes") if isinstance(context.get("nodes"), dict) else {}
    last = context.get("last") if isinstance(context.get("last"), dict) else {}
    return {**variables, "variables": variables, "nodes": nodes, "last": last}


def assign_path(target: dict[str, Any], key: str, value: Any) -> None:
    parts = key.split(".")
    current = target
    for part in parts[:-1]:
        nested = current.get(part)
        if not isinstance(nested, dict):
            nested = {}
            current[part] = nested
        current = nested
    current[parts[-1]] = value


def merge_usage(target: dict[str, Any], usage: dict[str, Any]) -> None:
    for key, value in usage.items():
        if isinstance(value, (int, float)):
            target[key] = target.get(key, 0) + value


def elapsed_ms(started_at: float) -> int:
    return int((time.perf_counter() - started_at) * 1000)


def maybe_decode_json_response(result: WorkflowLLMResult, node_data: dict[str, Any]) -> WorkflowLLMResult:
    response_format = node_data.get("response_format")
    if not isinstance(response_format, dict) or response_format.get("type") != "json_object":
        return result
    raw_content = result.answer
    if raw_content is None:
        return result
    if isinstance(raw_content, (dict, list)):
        return result
    if not isinstance(raw_content, str):
        return result
    decoded = decode_json_object_response(raw_content)
    if decoded is None:
        return result
    result.answer = decoded
    return result


def decode_json_object_response(raw_content: str) -> Any | None:
    fenced_candidates = [match.group(1).strip() for match in JSON_CODE_FENCE_PATTERN.finditer(raw_content)]
    candidates = [candidate for candidate in fenced_candidates if candidate]
    stripped = raw_content.strip()
    if stripped:
        candidates.append(stripped)
    for candidate in candidates:
        try:
            return json.loads(candidate)
        except json.JSONDecodeError:
            fragment = extract_json_fragment(candidate)
            if not fragment:
                continue
            try:
                return json.loads(fragment)
            except json.JSONDecodeError:
                continue
    return None


def extract_json_fragment(text: str) -> str | None:
    start_index: int | None = None
    expected_closings: list[str] = []
    in_string = False
    escaped = False
    for index, char in enumerate(text):
        if start_index is None:
            if char not in {"{", "["}:
                continue
            start_index = index
            expected_closings.append("}" if char == "{" else "]")
            continue
        if escaped:
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if char == '"':
            in_string = not in_string
            continue
        if in_string:
            continue
        if char in {"{", "["}:
            expected_closings.append("}" if char == "{" else "]")
            continue
        if char in {"}", "]"}:
            if not expected_closings or char != expected_closings[-1]:
                return None
            expected_closings.pop()
            if not expected_closings and start_index is not None:
                return text[start_index : index + 1]
    return None


def workflow_answer_text(value: Any) -> str:
    if isinstance(value, str):
        return value
    if value is None:
        return ""
    try:
        return json.dumps(value, ensure_ascii=False)
    except (TypeError, ValueError):
        return str(value)
